"""把 docs/ 下的 Markdown 交付文档转成 Word（.docx）

用法：
    python scripts/md_to_docx.py                     # 转 docs/ 下全部 .md
    python scripts/md_to_docx.py docs/05-测试报告.md  # 只转指定文件

为什么不用 Word COM：本机 Word 的 SaveAs2/Save 会无限挂起（新建文档写两个字就卡死）。
python-docx 直接构造 OOXML 包，不启动 Word，因此不受该问题影响。

支持的 Markdown 语法（覆盖交付文档实际用到的部分）：
    标题  # ## ### ####
    表格  | a | b |  （含 |---| 分隔行）
    代码块 ``` 围栏
    引用  >
    列表  - / * 与 1.
    分隔线 ---
    行内  **加粗**、`代码`
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

ROOT = Path(__file__).resolve().parent.parent

CN_BODY = "宋体"
CN_HEAD = "黑体"
EN_BODY = "Times New Roman"
MONO = "Consolas"

TABLE_SEP = re.compile(r"^\s*\|?[\s:|-]+\|[\s:|-]*$")
BULLET = re.compile(r"^\s*[-*]\s+(.*)$")
ORDERED = re.compile(r"^\s*(\d+)\.\s+(.*)$")
HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")
IMAGE = re.compile(r"^!\[(.*?)\]\((.+?)\)$")


def _set_font(run, cn=CN_BODY, en=EN_BODY, size=None, bold=None, color=None):
    """设置字体，并单独指定东亚字体——否则中文会退回默认字体"""
    run.font.name = en
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), en)
    rfonts.set(qn("w:hAnsi"), en)
    rfonts.set(qn("w:eastAsia"), cn)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def add_runs(par, text, base_size=10.5, mono=MONO):
    """写入一段文本，处理 **加粗** 与 `行内代码`"""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            r = par.add_run(piece[2:-2])
            _set_font(r, size=base_size, bold=True)
        elif piece.startswith("`") and piece.endswith("`"):
            r = par.add_run(piece[1:-1])
            _set_font(r, cn=MONO, en=MONO, size=base_size - 0.5)
            r.font.color.rgb = RGBColor(0xA8, 0x2D, 0x26)
        else:
            r = par.add_run(piece)
            _set_font(r, size=base_size)


def shade(cell_or_par, hex_fill):
    """给单元格或段落加底色"""
    el = cell_or_par._element
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    target = el.get_or_add_tcPr() if hasattr(el, "get_or_add_tcPr") else el.get_or_add_pPr()
    target.append(shd)


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = EN_BODY
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CN_BODY)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.4

    for name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12),
                       ("Heading 4", 11)):
        st = doc.styles[name]
        st.font.name = EN_BODY
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), CN_HEAD)
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(6)

    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(2.2)
        sec.left_margin = sec.right_margin = Cm(2.5)


def convert(md_path: Path, out_path: Path):
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    setup_styles(doc)

    i, in_code = 0, False
    code_buf = []

    while i < len(lines):
        line = lines[i]

        # ---- 代码块 ----
        if line.strip().startswith("```"):
            if in_code:
                par = doc.add_paragraph()
                par.paragraph_format.left_indent = Cm(0.5)
                par.paragraph_format.space_before = Pt(4)
                par.paragraph_format.space_after = Pt(8)
                par.paragraph_format.line_spacing = 1.15
                r = par.add_run("\n".join(code_buf))
                _set_font(r, cn=MONO, en=MONO, size=9)
                shade(par, "F5F5F5")
                code_buf, in_code = [], False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        stripped = line.strip()

        # ---- 空行 ----
        if not stripped:
            i += 1
            continue

        # ---- 分隔线 ----
        if re.fullmatch(r"-{3,}|\*{3,}", stripped):
            par = doc.add_paragraph()
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "BBBBBB")
            pbdr.append(bottom)
            par._element.get_or_add_pPr().append(pbdr)
            i += 1
            continue

        # ---- 表格 ----
        if stripped.startswith("|") and i + 1 < len(lines) and TABLE_SEP.match(lines[i + 1]):
            header = [c.strip() for c in stripped.strip("|").split("|")]
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            ncol = max([len(header)] + [len(r) for r in rows]) if rows else len(header)
            table = doc.add_table(rows=1, cols=ncol)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j in range(ncol):
                cell = table.rows[0].cells[j]
                cell.text = ""
                add_runs(cell.paragraphs[0], header[j] if j < len(header) else "",
                         base_size=9.5)
                for r in cell.paragraphs[0].runs:
                    r.font.bold = True
                shade(cell, "EDF1F5")
            for row in rows:
                cells = table.add_row().cells
                for j in range(ncol):
                    cells[j].text = ""
                    add_runs(cells[j].paragraphs[0],
                             row[j] if j < len(row) else "", base_size=9.5)
            doc.add_paragraph()
            continue

        # ---- 图片 ----
        m = IMAGE.match(stripped)
        if m:
            rel = m.group(2)
            img = (md_path.parent / rel)
            if img.exists():
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.add_run().add_picture(str(img), width=Cm(15.5))
                if m.group(1):
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r = cap.add_run(m.group(1))
                    _set_font(r, size=9)
                    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            else:
                print(f"  警告：图片不存在，已跳过 {img}")
            i += 1
            continue

        # ---- 标题 ----
        m = HEADING.match(stripped)
        if m:
            level = min(len(m.group(1)), 4)
            par = doc.add_heading(level=level)
            add_runs(par, m.group(2), base_size={1: 16, 2: 14, 3: 12, 4: 11}[level])
            for r in par.runs:
                r.font.bold = True
            i += 1
            continue

        # ---- 引用 ----
        if stripped.startswith(">"):
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.6)
            add_runs(par, stripped.lstrip("> ").strip(), base_size=10)
            for r in par.runs:
                r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1
            continue

        # ---- 列表 ----
        m = BULLET.match(line)
        if m:
            par = doc.add_paragraph(style="List Bullet")
            add_runs(par, m.group(1))
            i += 1
            continue
        m = ORDERED.match(line)
        if m:
            par = doc.add_paragraph(style="List Number")
            add_runs(par, m.group(2))
            i += 1
            continue

        # ---- 普通段落 ----
        par = doc.add_paragraph()
        add_runs(par, stripped)
        i += 1

    if in_code and code_buf:  # 收尾：围栏未闭合
        par = doc.add_paragraph()
        r = par.add_run("\n".join(code_buf))
        _set_font(r, cn=MONO, en=MONO, size=9)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def main():
    args = sys.argv[1:]
    if args:
        targets = [Path(a) if Path(a).is_absolute() else ROOT / a for a in args]
    else:
        targets = sorted((ROOT / "docs").glob("*.md"))
    if not targets:
        print("没有找到要转换的 Markdown 文件")
        return
    out_dir = ROOT / "docs" / "word"
    for md in targets:
        if not md.exists():
            print(f"跳过（不存在）：{md}")
            continue
        out = convert(md, out_dir / (md.stem + ".docx"))
        print(f"{md.name}  →  {out.relative_to(ROOT)}")
    print(f"\n共 {len(targets)} 份，输出目录：{out_dir}")


if __name__ == "__main__":
    main()
